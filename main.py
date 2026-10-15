from fastapi import FastAPI, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
import chardet
import hashlib
import json
import math
import pandas as pd
import os
import requests
import sqlite3
import tempfile
import time
import re
import observability
from pypdf import PdfReader
from docx import Document as DocxDocument
from pydantic import BaseModel
from dotenv import load_dotenv
from langchain_text_splitters import RecursiveCharacterTextSplitter
# 把嵌入模型缓存放进项目目录，而不是用户目录 ~/.cache/chroma
# （受限环境下用户目录不可写；放项目里也让整个项目自成一体）
from langchain_openai import ChatOpenAI
from langchain_core.tools import tool
from langchain.agents import create_agent
from langchain_core.messages import AIMessage
from langchain_chroma import Chroma             
from langchain_core.embeddings import Embeddings
from langchain_core.documents import Document
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.checkpoint.sqlite import SqliteSaver
import operator
from typing import Annotated, TypedDict
from langgraph.graph import StateGraph, START, END
from pathlib import Path
from chromadb.utils.embedding_functions.onnx_mini_lm_l6_v2 import ONNXMiniLM_L6_V2

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ONNXMiniLM_L6_V2.DOWNLOAD_PATH = (
    Path(BASE_DIR) / "models" / "onnx_models" / ONNXMiniLM_L6_V2.MODEL_NAME
)

# ==================== 基础配置 ====================
load_dotenv()

# ==================== 大模型配置（模型层可插拔） ====================
# 换模型 / 换供应商只需改 .env，代码不用动。兼容旧的 ARK_* 变量名。
LLM_BASE_URL = os.getenv("LLM_BASE_URL", "https://ark.cn-beijing.volces.com/api/v3").rstrip("/")
LLM_API_KEY = os.getenv("LLM_API_KEY") or os.getenv("ARK_API_KEY", "")
LLM_MODEL = os.getenv("LLM_MODEL") or os.getenv("ARK_MODEL", "ep-20261008111305-gwhdt")
# 排序/判断类任务可以换更轻量的模型来降本提速；不填则与主模型一致
RERANK_MODEL = os.getenv("RERANK_MODEL") or LLM_MODEL
# 关闭思考链的请求参数：各厂商字段名不同（通义是 enable_thinking），留空表示不发送；
# 若厂商不接受该参数，call_ai 会自动去掉后重试一次。
_NO_THINK_RAW = os.getenv("LLM_DISABLE_THINKING_JSON", '{"thinking": {"type": "disabled"}}')
try:
    NO_THINK_PARAMS = json.loads(_NO_THINK_RAW) if _NO_THINK_RAW.strip() else {}
except ValueError:
    NO_THINK_PARAMS = {}
app = FastAPI(title="AI智能应用平台")

# 跨域配置
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 上传目录与大小限制
UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "200"))
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024

# DataFrame 缓存：避免每次 /analyze 都重新解析大文件（uploads 里有 178MB 的 CSV）
DF_CACHE_MAX = int(os.getenv("DF_CACHE_SIZE", "3"))
_df_cache = {}  # abspath -> (mtime, size, DataFrame)

# ==================== 基础接口 ====================
@app.get("/")
def hello():
    return {"message": "后端启动成功！"}

# ==================== 文件上传接口（Agent模块用） ====================
@app.post("/upload")
def upload_file(file: UploadFile = File(...)):
    filename = safe_filename(file.filename)
    file_path = os.path.join(UPLOAD_DIR, filename)

    # 分块落盘：大文件不再一次性读进内存，超限立即中止
    size = 0
    too_large = False
    with open(file_path, "wb") as f:
        while True:
            chunk = file.file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                too_large = True
                break
            f.write(chunk)

    if too_large:
        try:
            os.remove(file_path)
        except OSError:
            pass
        return {"success": False, "error": f"文件超过 {MAX_UPLOAD_MB}MB 上限，请压缩或抽样后再上传"}

    try:
        df = load_dataframe(file_path)
    except Exception as e:
        return {"success": False, "error": f"文件解析失败：{str(e)[:120]}"}

    # 二进制/空文件常常不会抛异常，而是被解析成 0 行的空表，
    # 这种"成功"会让前端显示一份毫无意义的报告，直接拦掉
    if len(df) == 0 or len(df.columns) == 0:
        return {"success": False, "error": "文件里没有解析到任何数据行，请检查文件内容"}

    missing = df.isnull().sum().to_dict()
    missing_rate = (df.isnull().sum() / len(df) * 100).round(2).to_dict()

    return {
        "success": True,
        "filename": filename,
        "rows": len(df),
        "columns": len(df.columns),
        "column_names": df.columns.tolist(),
        "missing_count": to_native(missing),
        "missing_rate": to_native(missing_rate)
    }

# ==================== 工具函数 ====================
def safe_filename(filename):
    """只取文件名，挡住 ../../ 这类路径穿越"""
    return os.path.basename((filename or "").replace("\\", "/")).strip() or "upload.csv"


def read_csv_auto(file_path):
    """自动检测编码并读取CSV文件"""
    with open(file_path, 'rb') as f:
        raw_data = f.read(10000)
        detect_result = chardet.detect(raw_data)
        encoding = detect_result['encoding']

    if encoding and encoding.lower() in ['gb2312', 'gbk', 'cp936']:
        encoding = 'gbk'

    try:
        return pd.read_csv(file_path, encoding=encoding)
    except Exception:
        try:
            return pd.read_csv(file_path, encoding='utf-8')
        except Exception:
            return pd.read_csv(file_path, encoding='gbk')


def read_table_auto(file_path):
    """按后缀分流：CSV 走编码探测，Excel 走 openpyxl"""
    lower = file_path.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return pd.read_excel(file_path)
    if lower.endswith(".xls"):
        # 老式 .xls 需要额外的 xlrd，这里给出比栈信息更清楚的提示
        try:
            return pd.read_excel(file_path)
        except ImportError:
            raise ValueError("不支持老式 .xls 格式，请用 Excel 另存为 .xlsx 后重试")
    return read_csv_auto(file_path)


def load_dataframe(file_path):
    """带缓存地读取表格：同一文件（mtime + 大小未变）不重复解析"""
    key = os.path.abspath(file_path)
    stat = os.stat(file_path)
    cached = _df_cache.get(key)
    if cached and cached[0] == stat.st_mtime and cached[1] == stat.st_size:
        return cached[2]

    df = read_table_auto(file_path)

    if len(_df_cache) >= DF_CACHE_MAX:
        _df_cache.pop(next(iter(_df_cache)), None)
    _df_cache[key] = (stat.st_mtime, stat.st_size, df)
    return df


def read_text_auto(file_path):
    """按探测到的编码读文本，避免 GBK 文档被 errors='ignore' 吃掉中文"""
    with open(file_path, 'rb') as f:
        raw = f.read()
    encoding = chardet.detect(raw[:10000]).get('encoding') or 'utf-8'
    try:
        return raw.decode(encoding, errors='ignore')
    except (LookupError, UnicodeDecodeError):
        return raw.decode('utf-8', errors='ignore')

# 模块级缓存：ChatModel 只创建一次，复用底层 HTTP 连接池
_llm = None


def get_llm():
    """懒加载大模型客户端（LangChain ChatOpenAI，对接任意 OpenAI 兼容接口）"""
    global _llm
    if _llm is None:
        if not LLM_API_KEY:
            raise RuntimeError("未配置 LLM_API_KEY（或 ARK_API_KEY），请在项目根目录的 .env 中填写后重启后端")
        _llm = ChatOpenAI(
            model=LLM_MODEL,
            api_key=LLM_API_KEY,
            base_url=LLM_BASE_URL,
            temperature=0.1,
            timeout=120,
            max_retries=2,
        )
    return _llm


def call_ai(prompt: str, label: str = "chat", no_thinking: bool = False, model: str | None = None):
    """调用大模型（带耗时 / token / 成本埋点；支持任意 OpenAI 兼容接口）"""
    if not LLM_API_KEY:
        raise RuntimeError("未配置 LLM_API_KEY（或 ARK_API_KEY），请在项目根目录的 .env 中填写后重启后端")

    url = f"{LLM_BASE_URL}/chat/completions"

    headers = {
        "Authorization": f"Bearer {LLM_API_KEY}",
        "Content-Type": "application/json"
    }

    data = {
        "model": model or LLM_MODEL,
        "messages": [
            {"role": "system", "content": "你是一个数据分析专家，严格按照要求的格式输出。"},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.1
    }

    # 重排这类"判断/排序"任务不需要思考链：实测开启思考要 94.8s、烧掉 7900+ reasoning token，
    # 关闭后只需 2.5s，且答案完全一致。字段名因厂商而异，见 NO_THINK_PARAMS。
    extra = dict(NO_THINK_PARAMS) if no_thinking else {}
    data.update(extra)

    t0 = time.perf_counter()

    def elapsed():
        return (time.perf_counter() - t0) * 1000

    def post():
        """发一次请求；返回 (响应, 错误标记)"""
        try:
            # (连接超时, 读取超时)：上游不响应时不至于把请求永久挂住
            return requests.post(url, headers=headers, json=data, timeout=(10, 120)), None
        except requests.exceptions.Timeout:
            return None, "timeout"
        except requests.exceptions.RequestException as e:
            return None, str(e)

    response, err = post()
    # 有些厂商（尤其第三方模型）不认识"关闭思考"的参数，会返回 4xx —— 去掉该参数重试一次
    if response is not None and response.status_code >= 400 and extra:
        for k in extra:
            data.pop(k, None)
        response, err = post()

    if err == "timeout":
        observability.record_llm_call(label, elapsed(), ok=False, error="timeout")
        raise RuntimeError("调用大模型超时，请稍后重试")
    if err:
        observability.record_llm_call(label, elapsed(), ok=False, error=err)
        raise RuntimeError(f"调用大模型失败：{err[:120]}")

    if response.status_code != 200:
        observability.record_llm_call(label, elapsed(), ok=False, error=f"HTTP {response.status_code}")
        raise RuntimeError(f"大模型接口返回 {response.status_code}：{response.text[:200]}")

    try:
        result = response.json()
        usage = result.get("usage") or {}
        observability.record_llm_call(label, elapsed(),
                                      usage.get("prompt_tokens"), usage.get("completion_tokens"))
        return result["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError):
        observability.record_llm_call(label, elapsed(), ok=False, error="bad response")
        raise RuntimeError(f"大模型返回格式异常：{response.text[:200]}")


def to_native(data):
    """numpy类型转Python原生类型，解决JSON序列化问题"""
    if isinstance(data, dict):
        return {k: to_native(v) for k, v in data.items()}
    if isinstance(data, (list, tuple)):
        return [to_native(item) for item in data]
    # NaN / Inf 不是合法 JSON（会害前端 JSON.parse 直接抛错），统一转成 null
    if isinstance(data, float):
        return None if (math.isnan(data) or math.isinf(data)) else data
    if hasattr(data, 'item'):
        try:
            return to_native(data.item())
        except (ValueError, AttributeError):
            return data
    return data

# ==================== Agent工具库 ====================
def tool_get_overview(df):
    info = {
        "rows": len(df),
        "columns": len(df.columns),
        "column_names": df.columns.tolist(),
        "dtypes": df.dtypes.astype(str).to_dict()
    }
    info = to_native(info)
    text = f"数据概览：\n{json.dumps(info, ensure_ascii=False, indent=2)}"
    return text, info


def tool_get_stats(df, columns):
    result = to_native(df[columns].describe().round(2).to_dict())
    text = f"统计结果（列：{columns}）：\n{json.dumps(result, ensure_ascii=False, indent=2)}"
    return text, result


def tool_group_agg(df, group_col, agg_col, method="mean"):
    if method == "mean":
        res = df.groupby(group_col)[agg_col].mean().round(2)
    elif method == "sum":
        res = df.groupby(group_col)[agg_col].sum()
    elif method == "count":
        res = df.groupby(group_col)[agg_col].count()
    elif method == "max":
        res = df.groupby(group_col)[agg_col].max()
    elif method == "min":
        res = df.groupby(group_col)[agg_col].min()
    else:
        return "错误：不支持的聚合方法", None

    data = []
    for k, v in res.items():
        data.append({group_col: k, agg_col: v})
    data = to_native(data)

    text = f"分组聚合结果（按{group_col}统计{agg_col}的{method}）：\n{json.dumps(data, ensure_ascii=False, indent=2)}"
    return text, data


def tool_filter(df, column, operator, value):
    try:
        if operator == ">":
            filtered = df[df[column] > value]
        elif operator == "<":
            filtered = df[df[column] < value]
        elif operator == "==":
            filtered = df[df[column] == value]
        elif operator == ">=":
            filtered = df[df[column] >= value]
        elif operator == "<=":
            filtered = df[df[column] <= value]
        elif operator == "!=":
            filtered = df[df[column] != value]
        else:
            return "错误：不支持的运算符", None

        data = to_native(filtered.head(20).to_dict(orient="records"))
        text = f"过滤结果（{column} {operator} {value}）：共{len(filtered)}条数据\n前20行：\n{filtered.head(20).to_string()}"
        return text, data
    except Exception as e:
        return f"过滤执行失败：{str(e)}", None


def tool_sort(df, column, ascending=True):
    sorted_df = df.sort_values(by=column, ascending=ascending)
    data = to_native(sorted_df.head(20).to_dict(orient="records"))
    text = f"排序结果（按{column}，升序={ascending}）：\n前20行：\n{sorted_df.head(20).to_string()}"
    return text, data

def derive_chart(action, args, data):
    """从最后一次工具调用推导图表配置。

    比"让模型自己猜 x/y 字段"可靠：工具参数里本来就有真实的列名。
    返回 None 时前端会自动隐藏图表卡片。
    """
    if action == "group_agg":
        return {"type": "bar", "x_field": args.get("group_col"), "y_field": args.get("agg_col")}

    if action in ("filter", "sort") and isinstance(data, list) and data:
        keys = list(data[0].keys())
        x_field = keys[0]
        y_field = next(
            (k for k in keys
             if isinstance(data[0].get(k), (int, float)) and not isinstance(data[0].get(k), bool)),
            keys[-1],
        )
        return {"type": "bar", "x_field": x_field, "y_field": y_field}

    # get_overview / get_stats 是元信息和统计量，画图没有意义
    return None

# ==================== Agent数据分析接口 ====================
@app.post("/analyze")
def analyze_data(filename: str, question: str):
    file_path = os.path.join(UPLOAD_DIR, safe_filename(filename))
    if not os.path.exists(file_path):
        return {"success": False, "error": "文件不存在，请重新上传"}

    try:
        df = load_dataframe(file_path)
    except Exception as e:
        return {"success": False, "error": f"文件解析失败：{str(e)[:120]}"}

    captured = {"data": None}
    tools = build_tools(df, captured)

    agent = create_agent(
        model=get_llm(),
        tools=tools,
        system_prompt=(
            "你是专业的数据分析智能体，可以调用工具完成用户的数据分析任务。\n"
            "重要：每次调用工具之前，必须先用一句话（不超过 30 字）说明你接下来打算做什么、为什么，"
            "然后再调用工具。这句话会展示给用户看。\n"
            "1. 列名或数据类型不明确时，先调用 get_overview 了解数据结构\n"
            "2. 所有结论必须基于工具返回的真实结果，不要编造数据\n"
            "3. 任务完成后用中文给出简洁结论\n"
        ),
    )

    try:
        result = agent.invoke(
            {"messages": [{"role": "user", "content": question}]},
            # 可观测性回调：捕获 Agent 内部每一次模型调用的耗时与 token
            config={"callbacks": [observability.MetricsCallback("analyze")]},
        )
    except Exception as e:
        return {"success": False, "error": f"智能体执行失败：{str(e)[:200]}"}

    messages = result["messages"]

    # 把 LangGraph 的消息链路还原成前端需要的 tool_history
    tool_history = []
    last_action, last_input = None, {}
    for msg in messages:
        for call in (getattr(msg, "tool_calls", None) or []):
            tool_history.append({
                "step": len(tool_history) + 1,
                "thought": msg.content or "",
                "action": call["name"],
                "input": call["args"],
            })
            last_action, last_input = call["name"], call["args"]

    # 最后一条"有内容"的 AIMessage 就是最终结论
    final_answer = ""
    for msg in reversed(messages):
        if isinstance(msg, AIMessage) and msg.content:
            final_answer = msg.content
            break

    return {
        "success": True,
        "final_answer": final_answer,
        "chart": derive_chart(last_action, last_input, captured["data"]),
        "result": captured["data"],
        "tool_history": tool_history,
        "steps": len(tool_history),
    }
def build_tools(df, captured):
    """把工具绑定到当前 df，返回 LangChain Tool 列表。

    - df 不能交给模型，所以用闭包固定住，只暴露模型该填的参数
    - 工具返回值是给模型看的文本；结构化数据走 captured["data"] 旁路带出，供前端画图
    """

    @tool("get_overview", parse_docstring=True)
    def get_overview_tool() -> str:
        """获取数据概览：行数、列数、列名以及每列的数据类型。不需要参数。"""
        text, data = tool_get_overview(df)
        captured["data"] = data
        return text

    @tool("get_stats", parse_docstring=True)
    def get_stats_tool(columns: list[str]) -> str:
        """获取指定数值列的统计量（计数、均值、标准差、最小值、四分位数、最大值）。

        Args:
            columns: 要统计的列名列表，例如 ["2023年"]
        """
        text, data = tool_get_stats(df, columns)
        captured["data"] = data
        return text

    @tool("group_agg", parse_docstring=True)
    def group_agg_tool(group_col: str, agg_col: str, method: str = "mean") -> str:
        """按某列分组，对另一列做聚合计算。

        Args:
            group_col: 分组依据的列名
            agg_col: 被聚合的数值列名
            method: 聚合方法，可选 mean/sum/count/max/min
        """
        text, data = tool_group_agg(df, group_col, agg_col, method)
        captured["data"] = data
        return text

    @tool("filter", parse_docstring=True)
    def filter_tool(column: str, operator: str, value: str) -> str:
        """按条件过滤数据，返回符合条件的记录数与前 20 行。

        Args:
            column: 用于判断的列名
            operator: 运算符，可选 > < == >= <= !=
            value: 对比值
        """
        text, data = tool_filter(df, column, operator, value)
        captured["data"] = data
        return text

    @tool("sort", parse_docstring=True)
    def sort_tool(column: str, ascending: bool = True) -> str:
        """按指定列排序，返回前 20 行。

        Args:
            column: 排序依据的列名
            ascending: True 升序，False 降序
        """
        text, data = tool_sort(df, column, ascending)
        captured["data"] = data
        return text

    return [get_overview_tool, get_stats_tool, group_agg_tool, filter_tool, sort_tool]

CHROMA_PATH = "./chroma_db"
os.makedirs(CHROMA_PATH, exist_ok=True)


class ChromaONNXEmbeddings(Embeddings):
    """把 Chroma 自带的 ONNX 嵌入函数包装成 LangChain 的 Embeddings 接口。

    好处：直接复用已经下载到 models/ 的本地模型，不必再引入
    sentence-transformers + torch 那套重依赖。
    """

    def __init__(self):
        self._ef = ONNXMiniLM_L6_V2()

    def embed_documents(self, texts):
        return [[float(x) for x in v] for v in self._ef(list(texts))]

    def embed_query(self, text):
        return [float(x) for x in self._ef([text])[0]]


_kb_init_error = ""
try:
    vectorstore = Chroma(
        collection_name="knowledge_base",
        embedding_function=ChromaONNXEmbeddings(),
        persist_directory=CHROMA_PATH,
    )
except Exception as e:
    vectorstore = None
    _kb_init_error = str(e)

# 会话记忆：必须放在模块级，跨请求共享，否则每次请求都"失忆"
# 用 SQLite 持久化 —— 重启服务后对话记录依然在。
# check_same_thread=False 是必须的：FastAPI 的同步接口跑在线程池里，
# 不同请求可能落在不同线程，而 SQLite 连接默认只允许创建它的线程使用。
CHAT_DB_PATH = os.getenv("CHAT_DB_PATH") or os.path.join(BASE_DIR, "chat_memory.sqlite3")
try:
    _chat_conn = sqlite3.connect(CHAT_DB_PATH, check_same_thread=False)
    _rag_checkpointer = SqliteSaver(_chat_conn)
    _rag_checkpointer.setup()          # 建表（checkpoints / writes），幂等可重复调用
except Exception:
    # 持久化不可用时退回内存记忆，保证服务仍能启动
    _rag_checkpointer = InMemorySaver()


def kb_unavailable():
    return {"success": False, "error": f"知识库初始化失败：{_kb_init_error}"}


def split_text(text, chunk_size=500, chunk_overlap=50):
    """文本切片（LangChain RecursiveCharacterTextSplitter 实现）"""
    if not text:
        return []
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", "。", "！", "？", " ", ""],
    )
    return splitter.split_text(text)


# ==================== RAG接口：文档上传入库 ====================
@app.post("/rag/upload")
def rag_upload(file: UploadFile = File(...)):
    if vectorstore is None:
        return kb_unavailable()
    try:
        suffix = os.path.splitext(file.filename)[1].lower()
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(file.file.read())
            tmp_path = tmp.name

        # 解析文档文本
        text = ""
        if suffix == ".pdf":
            reader = PdfReader(tmp_path)
            for page in reader.pages:
                text += page.extract_text() + "\n\n"
        elif suffix == ".docx":
            doc = DocxDocument(tmp_path)      
            for para in doc.paragraphs:
                text += para.text + "\n"
        elif suffix in [".txt", ".md", ".csv"]:
            text = read_text_auto(tmp_path)
        else:
            return {"success": False, "error": f"不支持的文件格式：{suffix}"}

        # 切片并存库
        chunks = split_text(text)
        if not chunks:
            return {"success": False, "error": "文档中没有提取到可用文本（扫描版 PDF 常见）"}

        digest = hashlib.md5(text.encode("utf-8", errors="ignore")).hexdigest()[:8]
        ids = [f"{file.filename}-{digest}-{i}" for i in range(len(chunks))]
        vectorstore.add_documents(
            documents=[
                Document(page_content=c, metadata={"source": file.filename})
                for c in chunks
            ],
            ids=ids,
        )

        return {
            "success": True,
            "filename": file.filename,
            "chunks": len(chunks)
        }

    except Exception as e:
        return {"success": False, "error": str(e)}
    finally:
        if 'tmp_path' in dir():
            try:
                os.unlink(tmp_path)
            except:
                pass


# ==================== RAG接口：列出知识库文档 ====================
@app.get("/rag/documents")
def rag_documents():
    """列出知识库里已有的文档（按来源聚合片段数）"""
    if vectorstore is None:
        return kb_unavailable()
    try:
        data = vectorstore.get()
        counts = {}
        for meta in data.get("metadatas") or []:
            src = (meta or {}).get("source")
            if src:
                counts[src] = counts.get(src, 0) + 1
        docs = [{"source": s, "chunks": n} for s, n in sorted(counts.items())]
        return {
            "success": True,
            "documents": docs,
            "total_chunks": sum(counts.values()),
        }
    except Exception as e:
        return {"success": False, "error": str(e)}

# 两阶段检索参数：先宽召回，再由大模型重排
RECALL_K = 10    # 向量召回数量（候选越少，重排推理量越小）
FINAL_K = 3     # 重排后真正喂给模型的片段数



# 两阶段检索参数：先宽召回，再由大模型重排
_RERANK_PROMPT = """判断每段候选资料与用户提问的相关性，按相关性从高到低输出所有片段的编号。

只输出编号，用英文逗号分隔，不要解释、不要换行。
例如：2,7,0,5

用户提问：{question}

候选资料：
{snippets}
"""


def _parse_ranked(raw, n):
    """取出候选编号；兼容 '3,7,1' 与 {"ranked":[...]} 两种格式"""
    text = (raw or "").strip()
    if not text:
        return []
    start, end = text.find("{"), text.rfind("}")
    if start != -1 and end != -1:
        try:
            ranked = json.loads(text[start:end + 1]).get("ranked")
            if isinstance(ranked, list):
                return [i for i in ranked if isinstance(i, int) and 0 <= i < n]
        except Exception:
            pass
    # 退化处理：把输出里的整数当编号。
    # 注意正则要带 -?，否则 "-3" 会被 \d+ 抓成 "3"，负号丢失 → 静默取到错误片段
    nums = [int(x) for x in re.findall(r"-?\d+", text)]
    return [i for i in nums if 0 <= i < n][:n]

def rerank_documents(question, docs, top_n=FINAL_K):
    """第二阶段：用大模型对召回结果做相关性重排。

    任何异常（模型输出不合法、接口超时）都回退到向量排序，
    保证重排只可能变好、不会让检索挂掉。
    """
    docs = list(docs)
    if len(docs) <= 1:
        return docs[:top_n]
    snippets = "\n\n".join(f"[{i}] {d.page_content[:150]}" for i, d in enumerate(docs))
    try:
        raw = call_ai(_RERANK_PROMPT.format(question=question, snippets=snippets, top_n=top_n),
                      label="rerank", no_thinking=True, model=RERANK_MODEL)
        order = _parse_ranked(raw, len(docs))
        if not order:
            # 解析失败：记为失败指标，避免"表面正常、实际退化成向量排序"
            observability.record_llm_call("rerank-fallback", 0, ok=False, error="解析失败，回退向量排序")
            return docs[:top_n]
        return [docs[i] for i in order[:top_n]]
    except Exception as e:
        observability.record_llm_call("rerank-fallback", 0, ok=False, error=str(e)[:120])
        return docs[:top_n]

def build_list_documents_tool():
    """知识库的"元数据"工具：列出真实文档清单。

    为什么单独做一个工具：向量检索只能捞回 top-N 片段，模型会拿这几个片段
    反推"知识库一共几份文档"，从而过度概括（实测它把 6 份说成 2 份）。
    元数据类问题应当由工具直接给准确答案，而不是让检索去猜。
    """

    @tool("list_knowledge_documents", parse_docstring=True)
    def list_knowledge_documents() -> str:
        """列出知识库里已有的全部文档及其片段数。

        当用户询问"知识库里有哪些文档 / 有什么资料 / 一共几份文档"时使用本工具。
        """
        if vectorstore is None:
            return "知识库不可用。"
        data = vectorstore.get(include=["metadatas"])
        counts = {}
        for meta in data.get("metadatas") or []:
            src = (meta or {}).get("source")
            if src:
                counts[src] = counts.get(src, 0) + 1
        if not counts:
            return "知识库还是空的，还没有上传任何文档。"
        lines = [f"- {s}（{n} 个片段）" for s, n in sorted(counts.items())]
        return f"知识库中共有 {len(counts)} 份文档：\n" + "\n".join(lines)

    return list_knowledge_documents



def build_retriever_tool(captured, source=None, use_rerank=True):
    """构建 RAG 检索工具（两阶段：向量宽召回 + LLM 重排）。

    captured：把检索到的原文回传给接口层，用于溯源展示
    source：限定只在某一份文档里检索，避免"问 A 文档却答成 B 文档"
    """
    recall_kwargs = {"k": RECALL_K}
    if source:
        recall_kwargs["filter"] = {"source": source}
    wide_retriever = vectorstore.as_retriever(search_kwargs=recall_kwargs)

    @tool("search_knowledge_base", parse_docstring=True)
    def search_knowledge_base(query: str) -> str:
        """在本地知识库中检索与问题最相关的资料片段。

        Args:
            query: 检索用的自然语言问题
        """
        candidates = wide_retriever.invoke(query)
        docs = rerank_documents(query, candidates, top_n=FINAL_K) if use_rerank else candidates[:FINAL_K]
        captured["docs"] = [d.page_content for d in docs]
        if not docs:
            return "知识库中没有检索到相关内容。"
        return "\n\n".join(d.page_content for d in docs)

    return search_knowledge_base


class RagChatRequest(BaseModel):
    question: str
    session_id: str = "default"
    source: str | None = None
    # 是否启用 LLM 重排：精度更高，但每次多一次模型调用
    use_rerank: bool = True


_rag_system_prompt = (
    "你是知识库问答助手。\n"
    "1. 知识库里已经上传了文档。用户问“这份文档/这份文件/这个项目讲了什么”时，"
    "“这份/这个”指的就是知识库里的文档（若系统限定了检索范围，就是那份文档），"
    "必须直接调用 search_knowledge_base 检索，然后基于检索结果回答；不要反问“是哪份文档”。\n"
    "2. 用户问“知识库里有哪些文档 / 有什么资料 / 一共几份文档”时，"
    "必须调用 list_knowledge_documents 拿真实清单，**不要用检索到的几个片段反推文档总数**。\n"
    "3. 只要问题涉及知识库内容，都先检索，再严格基于检索结果回答，不要编造。\n"
    "4. 检索到的资料必须与问题主题相关才可使用。如果检索结果与用户问的主题**无关**"
    "（例如问的是 A 主题、检索到的却是 B 主题的资料），直接回答“知识库中没有与 XXX 相关的内容”，"
    "**绝不能用无关资料拼凑答案**。\n"
    "5. 只有用户明确询问对话本身（如“我刚才问的是什么”）时才不检索，直接结合上下文回答。\n"
    "6. 回答用中文。\n"
)


@app.post("/rag/chat")
def rag_chat(req: RagChatRequest):
    if vectorstore is None:
        return kb_unavailable()
    try:
        # 知识库为空时先给友好提示，省掉两次模型调用
        if vectorstore._collection.count() == 0:
            return {"success": False, "error": "知识库还是空的，请先上传文档"}

        # 检索工具：两阶段检索（宽召回 + 重排），按需限定文档范围
        collected = {"docs": []}
        search_tool = build_retriever_tool(collected, source=req.source, use_rerank=req.use_rerank)

        agent = create_agent(
            model=get_llm(),
            tools=[search_tool, build_list_documents_tool()],
            checkpointer=_rag_checkpointer,
            system_prompt=_rag_system_prompt,
        )

        result = agent.invoke(
            {"messages": [{"role": "user", "content": req.question}]},
            config={
                # thread_id 就是会话 ID：同一个 session_id 的多次提问共享上下文
                "configurable": {"thread_id": req.session_id},
                # 可观测性回调
                "callbacks": [observability.MetricsCallback("rag")],
            },
        )

        return {
            "success": True,
            "answer": result["messages"][-1].content,
            "related_docs": collected["docs"],
        }
    except Exception as e:
        return {"success": False, "error": str(e)}

        # ==================== 可观测性：指标查询 ====================
@app.get("/metrics")
def get_metrics():
    """大模型调用指标：次数 / token / 延迟（均值、P95）/ 估算成本"""
    return observability.snapshot()


# ==================== 多智能体协同（Supervisor 架构） ====================
# 规划节点判断需要哪些专家 → 专家各自执行 → 汇总节点融合结论。
# 与单 Agent 的区别：单个问题可以同时调动"数据分析"和"知识库检索"两个专家，
# 例如"文档里的结论，和我上传的这份数据对得上吗？"

_MULTI_PLAN_PROMPT = """你是一个任务调度器。根据用户问题，判断需要哪些专家来回答。

可选专家：
- data：数据分析专家。能读取用户上传的表格，做统计、分组、过滤、排序等真实计算。
- kb：知识库专家。能检索知识库中的文档（项目分析报告等），回答"文档里写了什么"。

规则：
1. 需要基于表格算出真实结果 → data
2. 问文档内容、文档结论 → kb
3. 需要把两者结合（如"文档的结论和我的数据对得上吗"）→ 两个都选
4. 只输出 JSON，不要解释。格式：{{"agents": ["data"]}} 或 {{"agents": ["kb"]}} 或 {{"agents": ["data", "kb"]}}

用户问题：{question}
{file_hint}
"""

_MULTI_SUMMARY_PROMPT = """你是多智能体系统的汇总器。下面是各专家针对同一问题给出的独立结论，
请把它们融合成一段完整、简洁的中文回答。

要求：
1. 保留关键数字与事实，不得编造任何信息
2. 如果各专家的结论存在**矛盾或不一致**，必须明确指出
3. 不要罗列"专家A说…专家B说…"，直接给出融合后的结论

用户问题：{question}

各专家结论：
{findings}
"""


class MultiState(TypedDict):
    question: str
    filename: str | None
    plan: list
    findings: Annotated[list, operator.add]
    answer: str


def _multi_plan(state):
    """① 规划：决定调用哪些专家（分类任务，关闭思考更快）"""
    has_file = bool(state.get("filename"))
    hint = f"（用户已上传表格：{state['filename']}）" if has_file else "（用户未上传表格，data 专家不可用）"
    try:
        raw = call_ai(_MULTI_PLAN_PROMPT.format(question=state["question"], file_hint=hint),
                      label="multi/plan", no_thinking=True)
        s, e = raw.find("{"), raw.rfind("}")
        agents = json.loads(raw[s:e + 1]).get("agents") or []
    except Exception:
        agents = []

    agents = [a for a in agents if a in ("data", "kb")]
    if not has_file:                     # 没有表格，数据专家无法工作
        agents = [a for a in agents if a != "data"]
    if vectorstore is None:              # 知识库不可用
        agents = [a for a in agents if a != "kb"]
    if not agents:                       # 兜底：按可用性选一个
        agents = ["data"] if has_file else (["kb"] if vectorstore is not None else [])
    return {"plan": agents}


def _multi_data_expert(state):
    """② 专家一：数据分析（复用现有 5 个工具）"""
    t0 = time.perf_counter()
    file_path = os.path.join(UPLOAD_DIR, safe_filename(state["filename"]))
    df = load_dataframe(file_path)
    captured = {"data": None}
    agent = create_agent(
        model=get_llm(),
        tools=build_tools(df, captured),
        system_prompt=(
            "你是数据分析专家。所有结论必须基于工具返回的真实结果，不要编造数据。"
            "遇到不认识的字段编码（如 FIPS 州代码、行业代码）不要臆测其含义，只描述数值本身。"
            "用中文给出简洁结论。"
        ),
    )
    result = agent.invoke(
        {"messages": [{"role": "user", "content": state["question"]}]},
        config={"callbacks": [observability.MetricsCallback("multi/data")]},
    )
    answer = ""
    for msg in reversed(result["messages"]):
        if isinstance(msg, AIMessage) and msg.content:
            answer = msg.content
            break
    return {"findings": [{
        "agent": "data", "title": "数据分析专家", "content": answer,
        "elapsed_ms": round((time.perf_counter() - t0) * 1000),
        "file": state["filename"],
    }]}


def _multi_kb_expert(state):
    """③ 专家二：知识库检索（复用 RAG 两阶段检索）"""
    t0 = time.perf_counter()
    captured = {"docs": []}
    agent = create_agent(
        model=get_llm(),
        tools=[build_retriever_tool(captured), build_list_documents_tool()],
        system_prompt=_rag_system_prompt,
    )
    result = agent.invoke(
        {"messages": [{"role": "user", "content": state["question"]}]},
        config={"callbacks": [observability.MetricsCallback("multi/kb")]},
    )
    return {"findings": [{
        "agent": "kb", "title": "知识库专家",
        "content": result["messages"][-1].content,
        "elapsed_ms": round((time.perf_counter() - t0) * 1000),
        "docs": captured["docs"],
    }]}


def _multi_summarize(state):
    """④ 汇总：融合各专家结论（只有一位专家时省掉这次调用）"""
    findings = state.get("findings") or []
    if not findings:
        return {"answer": "抱歉，没有可用的资源来回答这个问题。"}
    if len(findings) == 1:
        return {"answer": findings[0]["content"]}
    blocks = "\n\n".join(f"【{f['title']}】\n{f['content']}" for f in findings)
    answer = call_ai(_MULTI_SUMMARY_PROMPT.format(question=state["question"], findings=blocks),
                     label="multi/summary")
    return {"answer": answer}


def _route_after_plan(state):
    agents = state.get("plan") or []
    if "data" in agents:
        return "data_expert"
    if "kb" in agents:
        return "kb_expert"
    return "summarize"


def _route_after_data(state):
    return "kb_expert" if "kb" in (state.get("plan") or []) else "summarize"


def build_multi_agent():
    """构建 Supervisor 图：规划 → 专家（串行，可多个）→ 汇总"""
    g = StateGraph(MultiState)
    g.add_node("plan", _multi_plan)
    g.add_node("data_expert", _multi_data_expert)
    g.add_node("kb_expert", _multi_kb_expert)
    g.add_node("summarize", _multi_summarize)

    g.add_edge(START, "plan")
    g.add_conditional_edges("plan", _route_after_plan, ["data_expert", "kb_expert", "summarize"])
    g.add_conditional_edges("data_expert", _route_after_data, ["kb_expert", "summarize"])
    g.add_edge("kb_expert", "summarize")
    g.add_edge("summarize", END)
    return g.compile()


_multi_agent = None


def get_multi_agent():
    global _multi_agent
    if _multi_agent is None:
        _multi_agent = build_multi_agent()
    return _multi_agent


class MultiRequest(BaseModel):
    question: str
    filename: str | None = None


@app.post("/multi/analyze")
def multi_analyze(req: MultiRequest):
    """多智能体协同：规划 → 专家执行 → 汇总"""
    if req.filename:
        file_path = os.path.join(UPLOAD_DIR, safe_filename(req.filename))
        if not os.path.exists(file_path):
            return {"success": False, "error": "文件不存在，请重新上传"}
    if vectorstore is None and not req.filename:
        return {"success": False, "error": "知识库不可用，且没有上传表格文件，无可用的专家"}

    t0 = time.perf_counter()
    try:
        state = get_multi_agent().invoke({
            "question": req.question,
            "filename": req.filename,
            "plan": [],
            "findings": [],
            "answer": "",
        })
    except Exception as e:
        return {"success": False, "error": f"多智能体执行失败：{str(e)[:200]}"}

    return {
        "success": True,
        "plan": state.get("plan") or [],
        "findings": state.get("findings") or [],
        "answer": state.get("answer") or "",
        "elapsed_ms": round((time.perf_counter() - t0) * 1000),
    }